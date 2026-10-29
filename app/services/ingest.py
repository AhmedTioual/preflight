"""Read a client file (Excel or CSV) into headers and rows of raw values."""

import csv
import hashlib
import io
from dataclasses import dataclass
from datetime import date, datetime

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
ALLOWED = (".xlsx", ".csv")


class IngestError(ValueError):
    pass


@dataclass
class ParsedFile:
    filename: str
    sha256: str
    columns: list[str]
    rows: list[dict]


def _clean(value):
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _table(records: list[list]) -> tuple[list[str], list[dict]]:
    records = [r for r in records if any(_clean(v) is not None for v in r)]
    if not records:
        raise IngestError("The file is empty.")
    header_row = records[0]
    columns, seen = [], set()
    for i, h in enumerate(header_row):
        name = str(_clean(h) or f"Column {i + 1}")
        while name in seen:
            name = f"{name} ({i + 1})"
        seen.add(name)
        columns.append(name)
    rows = []
    for record in records[1:]:
        values = list(record) + [None] * (len(columns) - len(record))
        row = {col: _clean(v) for col, v in zip(columns, values, strict=False)}
        if any(v is not None for v in row.values()):
            rows.append(row)
    if not rows:
        raise IngestError("The file has a header row but no data rows.")
    if len(rows) > MAX_ROWS:
        raise IngestError(f"The file has {len(rows)} rows; the limit is {MAX_ROWS}.")
    return columns, rows


def _read_xlsx(data: bytes) -> list[list]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # corrupt or not really an xlsx
        raise IngestError("The Excel file could not be opened.") from exc
    ws = wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _read_csv(data: bytes) -> list[list]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise IngestError("The CSV file encoding is not supported.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)]


def read_file(filename: str, data: bytes) -> ParsedFile:
    name = filename.lower()
    if not name.endswith(ALLOWED):
        raise IngestError("Upload an Excel (.xlsx) or CSV file.")
    if len(data) > MAX_BYTES:
        raise IngestError("The file is larger than 5 MB.")
    records = _read_xlsx(data) if name.endswith(".xlsx") else _read_csv(data)
    columns, rows = _table(records)
    return ParsedFile(filename, hashlib.sha256(data).hexdigest(), columns, rows)
