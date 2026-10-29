"""Generate the synthetic client files in sample_files/.

Usage:  python -m scripts.make_samples
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app import demo_data as d

OUT = Path(__file__).resolve().parents[1] / "sample_files"


def write_xlsx(path: Path, headers: list[str], rows: list[list]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Fund data"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
    for row in rows:
        ws.append(row)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(14, len(str(col[0].value)) + 2)
    wb.save(path)


def write_csv(path: Path, headers: list[str], rows: list[list]) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(headers)
        writer.writerows(rows)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    write_xlsx(OUT / "northstar_fund_return_2026-06-30.xlsx", d.NS_HEADERS, d.NS_Q2_V1)
    write_xlsx(OUT / "northstar_fund_return_2026-06-30_v2.xlsx", d.NS_HEADERS, d.NS_Q2_V2)
    write_xlsx(OUT / "northstar_fund_return_2026-09-30.xlsx", d.NS_HEADERS, d.NS_Q3)
    write_csv(OUT / "riverbend_declaration_2026-03-31.csv", d.RB_HEADERS, d.RB_Q1)
    write_csv(OUT / "riverbend_declaration_2026-06-30.csv", d.RB_HEADERS, d.RB_Q2)
    print(f"Sample files written to {OUT}")


if __name__ == "__main__":
    main()
